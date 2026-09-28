"""Üçel teknik veri formu — Excel sürümü (telefondan / bilgisayardan doldurmak için).

Alanlar teknik-veri.mjs ve katalog2-veri.mjs'ten okunur; PDF form ve katalog
tablolarıyla birebir aynıdır.
Çalıştırma (bu klasörde):  python3 form_excel.py
Çıktı: cikti/Ucel-Teknik-Veri-Formu.xlsx
"""
import json
import subprocess
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

BURASI = Path(__file__).resolve().parent
CIKTI = BURASI / 'cikti'

JS = """
Promise.all([import('./teknik-veri.mjs'), import('./katalog2-veri.mjs')]).then(([t, k]) => {
  process.stdout.write(JSON.stringify({ firma: k.FIRMA, urunler: k.URUNLER.map((u) => ({ id: u.id, ad: u.ad })),
    teknik: t.TEKNIK, diger: t.DIGER_TEKNIK, standart: t.STANDART_CEKIM }));
});
"""
# node -e (CommonJS) içinden dinamik import ile ESM veri dosyaları okunur
V = json.loads(subprocess.run(['node', '-e', JS], cwd=BURASI, capture_output=True, text=True, check=True).stdout)

FONT = 'Arial'
KOMUR, ALTIN, GRI = '1B1B1D', 'A9824F', '6B665F'
SARI = PatternFill('solid', fgColor='FFF2CC')      # doldurulacak hücre
BASLIK_DOLGU = PatternFill('solid', fgColor='1B1B1D')
ACIK = PatternFill('solid', fgColor='F2F0EC')
INCE = Side(style='thin', color='C9C2B6')
KENAR = Border(left=INCE, right=INCE, top=INCE, bottom=INCE)
SARMA = Alignment(wrap_text=True, vertical='center')
MODEL = 3


def yaz(ws, hucre, deger, *, kalin=False, renk=KOMUR, boyut=10, dolgu=None, italik=False, kenar=False, hiza=SARMA):
    c = ws[hucre]
    c.value = deger
    c.font = Font(name=FONT, bold=kalin, color=renk, size=boyut, italic=italik)
    c.alignment = hiza
    if dolgu:
        c.fill = dolgu
    if kenar:
        c.border = KENAR
    return c


def baslik_satiri(ws, satir, basliklar):
    for i, b in enumerate(basliklar):
        c = ws.cell(row=satir, column=i + 1, value=b)
        c.font = Font(name=FONT, bold=True, color='FFFFFF', size=9)
        c.fill = BASLIK_DOLGU
        c.alignment = SARMA
        c.border = KENAR


def girdi(ws, satir, sutun, deger=None):
    c = ws.cell(row=satir, column=sutun, value=deger)
    c.fill = SARI
    c.border = KENAR
    c.font = Font(name=FONT, size=10, color=KOMUR)
    c.alignment = SARMA
    return c


def etiket(ws, satir, sutun, deger, *, renk=KOMUR, kalin=False, boyut=10):
    c = ws.cell(row=satir, column=sutun, value=deger)
    c.border = KENAR
    c.font = Font(name=FONT, size=boyut, color=renk, bold=kalin)
    c.alignment = SARMA
    return c


def liste(ws, secenekler, aralik):
    dv = DataValidation(type='list', formula1='"' + ','.join(secenekler) + '"', allow_blank=True)
    dv.error = 'Listeden seçin ya da boş bırakın.'
    dv.errorTitle = 'Geçersiz seçim'
    ws.add_data_validation(dv)
    dv.add(aralik)


def sayfa_basligi(ws, baslik, aciklama):
    yaz(ws, 'A1', 'ÜÇEL TARIM ALETLERİ · TEKNİK VERİ FORMU', kalin=True, renk=ALTIN, boyut=9)
    yaz(ws, 'A2', baslik, kalin=True, boyut=16)
    yaz(ws, 'A3', aciklama, renk=GRI, boyut=9, italik=True)
    ws.row_dimensions[2].height = 26


wb = Workbook()

# ---------- 1. Nasıl doldurulur ----------
ws = wb.active
ws.title = 'Nasıl Doldurulur'
ws.sheet_view.showGridLines = False
ws.column_dimensions['A'].width = 30
ws.column_dimensions['B'].width = 62
sayfa_basligi(ws, 'Nasıl doldurulur?', 'Bu dosyadaki bilgiler katalogdaki teknik tablolara birebir aktarılacak.')
kurallar = [
    ('Sarı hücreler', 'Yalnızca sarı hücrelere yazın. Diğer hücreler açıklamadır.'),
    ('Bilmediğiniz alan', 'Boş bırakın. Tahmin yazmayın. Boş alanlar katalogda gösterilmez.'),
    ('Birden fazla model', 'Her model için ayrı sütun kullanın (Model 1, 2, 3). Tek model varsa yalnızca Model 1.'),
    ('Birimler', 'Birim sütunundaki birimle yazın: ton, kg, mm, cm, lt, HP.'),
    ('Açılır listeler', 'Evet/Hayır, Var/Yok gibi alanlarda hücreye dokunup listeden seçin.'),
    ('Göndermek', 'Dosyayı kaydedip formu size ileten kişiye WhatsApp ya da e-postayla gönderin.'),
]
for i, (a, b) in enumerate(kurallar, start=5):
    etiket(ws, i, 1, a, kalin=True)
    etiket(ws, i, 2, b)
    ws.row_dimensions[i].height = 30
yaz(ws, 'A12', 'Renk anlamı', kalin=True)
girdi(ws, 13, 1, '')
etiket(ws, 13, 2, 'Sarı hücre = sizin dolduracağınız alan')

# Örnek satır (biçim göstermek için; gerçek veri değildir)
yaz(ws, 'A15', 'ÖRNEK — yazım biçimi (gerçek veri değildir, dikkate alınmaz)', kalin=True, renk=GRI)
baslik_satiri(ws, 16, ['Özellik', 'Model 1 sütununa yazılış örneği'])
for i, (a, b) in enumerate([('Model adı / kodu', 'Örnek-01'), ('Kasa iç ölçüsü (U × G × Y)', '3000 × 2000 × 400 mm'), ('Gereken traktör gücü', '50–70 HP'), ('Damper', 'arkaya')], start=17):
    etiket(ws, i, 1, a, renk=GRI)
    etiket(ws, i, 2, b, renk=GRI).font = Font(name=FONT, size=10, color=GRI, italic=True)

# Doluluk özeti (formüllerle)
yaz(ws, 'A23', 'Doluluk durumu (otomatik)', kalin=True)
baslik_satiri(ws, 24, ['Sayfa', 'Doldurulan teknik hücre'])
ozet_satir = 25

# ---------- 2. Firma ----------
wf = wb.create_sheet('Firma')
wf.sheet_view.showGridLines = False
for col, w in zip('ABCD', (30, 40, 14, 40)):
    wf.column_dimensions[col].width = w
sayfa_basligi(wf, 'Firma bilgileri', 'Mevcut bilgiyi kontrol edin: doğruysa "Evet", değilse "Hayır" seçip düzeltmeyi yazın.')
baslik_satiri(wf, 5, ['Bilgi', 'Sitede yazan', 'Doğru mu?', 'Düzeltme / cevap'])
F = V['firma']
firma_satirlari = [
    ('Katalogda kullanılacak ad', 'Üçel Tarım Aletleri (seçenekler: Üçel Zirai Aletler, Üçel Ziraat)'),
    ('Resmî / ticari unvan', ''),
    ('Adres', f"{F['adres1']}, {F['adres2']}"),
    ('Telefon / WhatsApp', F['telefon']),
    ('Çalışma saatleri', F['saatler']),
    ('Yetkili', F['yetkili']),
    ('Kaç yıldır üretim yapılıyor? (yalnızca belgeliyse kullanılır)', ''),
    ('İmalat ürünlerinde ortalama teslim süresi', ''),
    ('Nakliye nasıl yapılıyor?', ''),
    ('Katalogdaki müşteri fotoğrafları için izin', ''),
]
for i, (a, b) in enumerate(firma_satirlari, start=6):
    etiket(wf, i, 1, a, kalin=True)
    etiket(wf, i, 2, b, renk=GRI)
    if b:
        girdi(wf, i, 3)
    else:
        etiket(wf, i, 3, '—', renk=GRI)
    girdi(wf, i, 4)
    wf.row_dimensions[i].height = 32
liste(wf, ['Evet', 'Hayır'], 'C6:C15')
dv_nak = DataValidation(type='list', formula1='"Biz götürüyoruz,Anlaşmalı firma,Müşteri teslim alıyor"', allow_blank=True)
wf.add_data_validation(dv_nak); dv_nak.add('D14')
dv_izin = DataValidation(type='list', formula1='"İzin alındı,İzin alınacak,Yüzler bulanıklaştırılsın"', allow_blank=True)
wf.add_data_validation(dv_izin); dv_izin.add('D15')
wf.freeze_panes = 'A6'

# ---------- 3. Ürün listesi ----------
wl = wb.create_sheet('Ürün Listesi')
wl.sheet_view.showGridLines = False
for col, w in zip('ABCDE', (38, 16, 18, 12, 40)):
    wl.column_dimensions[col].width = w
sayfa_basligi(wl, 'Ürün listesi', '"Kendi imalat" yalnızca Göksun atölyesinde üretilen ürünler içindir; katalogdaki rozet buna göre konur.')
baslik_satiri(wl, 5, ['Ürün', 'Satıyor musunuz?', 'Kimin üretimi?', 'Model sayısı', 'Model adları'])
urun_adlari = [u['ad'] for u in V['urunler']] + ['Mibzer', 'Çayır Biçme Makinesi', 'Döner Ot Tırmığı', 'Tesviye Küreği',
                                                 'Kırmızı hidrolik ekipman (fotoğraflarda) — adını E sütununa yazın', 'Diğer:', 'Diğer:']
for i, a in enumerate(urun_adlari, start=6):
    if a == 'Diğer:':
        girdi(wl, i, 1, None)
    else:
        etiket(wl, i, 1, a, kalin=True)
    for c in range(2, 6):
        girdi(wl, i, c)
    wl.row_dimensions[i].height = 22
son = 5 + len(urun_adlari)
liste(wl, ['Evet', 'Hayır'], f'B6:B{son}')
liste(wl, ['Kendi imalat', 'Satış'], f'C6:C{son}')
dv_n = DataValidation(type='whole', operator='between', formula1='0', formula2='20', allow_blank=True)
dv_n.error = '0 ile 20 arasında bir sayı yazın.'
wl.add_data_validation(dv_n); dv_n.add(f'D6:D{son}')
etiket(wl, son + 2, 1, 'Yedek parça hangi ürünler için var?', kalin=True)
girdi(wl, son + 2, 2); wl.merge_cells(start_row=son + 2, start_column=2, end_row=son + 2, end_column=5)
wl.row_dimensions[son + 2].height = 40
wl.freeze_panes = 'A6'


# ---------- 4. Ürün sayfaları ----------
def urun_sayfasi(ad, teknik, uyum=None, opsiyon=None, cekim=None, sayfa_adi=None):
    w = wb.create_sheet(sayfa_adi or ad[:31])
    w.sheet_view.showGridLines = False
    for col, wd in zip('ABCDE', (36, 22, 20, 20, 20)):
        w.column_dimensions[col].width = wd
    sayfa_basligi(w, ad, 'Her model için bir sütun. Bilmediğiniz hücreyi boş bırakın.')
    r = 5
    yaz(w, f'A{r}', '1. TEKNİK ÖZELLİKLER', kalin=True, renk=ALTIN, boyut=9)
    r += 1
    baslik_satiri(w, r, ['Özellik', 'Birim / seçenek'] + [f'Model {i + 1}' for i in range(MODEL)])
    ilk = r + 1
    satirlar = [('Model adı / kodu', '')] + [(a, b) for a, b, *_ in teknik if a != 'Model']
    for i, (a, b) in enumerate(satirlar):
        rr = ilk + i
        etiket(w, rr, 1, a, kalin=(i == 0))
        etiket(w, rr, 2, b, renk=GRI, boyut=9)
        for c in range(3, 3 + MODEL):
            girdi(w, rr, c)
        w.row_dimensions[rr].height = 20
    son_t = ilk + len(satirlar) - 1
    teknik_aralik = f"C{ilk}:E{son_t}"
    r = son_t + 2
    if uyum:
        yaz(w, f'A{r}', '2. UYUM', kalin=True, renk=ALTIN, boyut=9); r += 1
        baslik_satiri(w, r, ['Kriter', 'Cevap']); w.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5); r += 1
        for k in uyum:
            etiket(w, r, 1, k[0]); girdi(w, r, 2)
            w.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
            w.row_dimensions[r].height = 20; r += 1
        r += 1
    if opsiyon:
        yaz(w, f'A{r}', '3. OPSİYONLAR', kalin=True, renk=ALTIN, boyut=9); r += 1
        baslik_satiri(w, r, ['Opsiyon', 'Var mı?', 'Açıklama']); w.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5); r += 1
        o_ilk = r
        for o in opsiyon + [['Diğer (yazın):', None]]:
            if o[0] == 'Diğer (yazın):':
                girdi(w, r, 1, None)
            else:
                etiket(w, r, 1, o[0])
            girdi(w, r, 2); girdi(w, r, 3)
            w.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
            w.row_dimensions[r].height = 20; r += 1
        liste(w, ['Var', 'Yok'], f'B{o_ilk}:B{r - 1}')
        r += 1
    if cekim is not None:
        yaz(w, f'A{r}', '4. FOTOĞRAF (çekildikçe "Çekildi" seçin)', kalin=True, renk=ALTIN, boyut=9); r += 1
        baslik_satiri(w, r, ['Kare', 'Durum']); r += 1
        f_ilk = r
        for k in V['standart'] + cekim:
            etiket(w, r, 1, k); girdi(w, r, 2)
            w.row_dimensions[r].height = 30 if len(k) > 40 else 20; r += 1
        liste(w, ['Çekildi', 'Çekilecek', 'Mümkün değil'], f'B{f_ilk}:B{r - 1}')
        r += 1
    yaz(w, f'A{r}', 'NOT', kalin=True, renk=ALTIN, boyut=9); r += 1
    girdi(w, r, 1); w.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5); w.row_dimensions[r].height = 60
    w.freeze_panes = f'C{ilk}'
    return w.title, teknik_aralik


ozet = []
adlar = {u['id']: u['ad'] for u in V['urunler']}
for uid, t in V['teknik'].items():
    sayfa_adi = adlar[uid].replace('/', '-').replace('  ', ' ')[:31]
    ozet.append(urun_sayfasi(adlar[uid], t['teknik'], t['uyum'], t['opsiyon'], t['cekim'], sayfa_adi))

for ad, alanlar in V['diger'].items():
    teknik = []
    for a in alanlar:
        if '(' in a and a.endswith(')'):
            b, birim = a[:-1].split(' (', 1)
            teknik.append([b, birim, None])
        else:
            teknik.append([a, '', None])
    teknik.insert(1, ['Marka', '', None])
    ozet.append(urun_sayfasi(ad, teknik, cekim=[], sayfa_adi=ad[:31]))

# ---------- 5. Onay ----------
wo = wb.create_sheet('Onay')
wo.sheet_view.showGridLines = False
wo.column_dimensions['A'].width = 28; wo.column_dimensions['B'].width = 60
sayfa_basligi(wo, 'Onay', 'Bu dosyadaki bilgiler katalogda ve internet sitesinde kullanılacaktır. Boş alanlar gösterilmez.')
for i, a in enumerate(['Ad Soyad', 'Tarih', 'Eklemek istedikleriniz'], start=5):
    etiket(wo, i, 1, a, kalin=True); girdi(wo, i, 2)
wo.row_dimensions[7].height = 120

# Doluluk özeti formülleri
for i, (sayfa, aralik) in enumerate(ozet):
    r = ozet_satir + i
    etiket(ws, r, 1, sayfa)
    c = ws.cell(row=r, column=2, value=f"=COUNTA('{sayfa}'!{aralik})")
    c.font = Font(name=FONT, size=10, color=KOMUR); c.border = KENAR
etiket(ws, ozet_satir + len(ozet), 1, 'Toplam', kalin=True)
t = ws.cell(row=ozet_satir + len(ozet), column=2, value=f'=SUM(B{ozet_satir}:B{ozet_satir + len(ozet) - 1})')
t.font = Font(name=FONT, size=10, bold=True); t.border = KENAR

for s in wb.worksheets:
    s.sheet_properties.tabColor = ALTIN if s.title in ('Nasıl Doldurulur', 'Firma', 'Ürün Listesi', 'Onay') else 'D2B181'
    s.page_setup.orientation = 'portrait'
    s.page_setup.fitToWidth = 1
    s.page_setup.fitToHeight = 0
    s.sheet_properties.pageSetUpPr.fitToPage = True

CIKTI.mkdir(exist_ok=True)
hedef = CIKTI / 'Ucel-Teknik-Veri-Formu.xlsx'
wb.save(hedef)
print('yazıldı:', hedef, '·', len(wb.worksheets), 'sayfa')
